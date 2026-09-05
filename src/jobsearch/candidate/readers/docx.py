from __future__ import annotations

from pathlib import Path

from .base import DocumentReader, ReaderError


class DocxReader(DocumentReader):
    extensions = (".docx",)

    def read(self, path: Path) -> str:
        try:
            import docx
        except ImportError as exc:
            raise ReaderError("python-docx is required to read .docx files") from exc
        try:
            document = docx.Document(str(path))
        except Exception as exc:
            raise ReaderError(f"could not read {path}: {exc}") from exc
        parts = [p.text for p in document.paragraphs]
        # CVs very often keep skills and dates in tables, and paragraph
        # iteration alone silently skips every one of them.
        for table in document.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        return "\n".join(p for p in parts if p.strip()).strip()
